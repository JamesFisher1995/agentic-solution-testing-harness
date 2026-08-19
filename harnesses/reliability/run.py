"""Reliability harness — asserts the AI solution recommends the right tests to catch a known defect.

Design intent
-------------
For each defined defect scenario we submit the related CR to the solution N times and
assert that the test(s) which would catch the defect are recommended as Must Run or
Should Run.  If the solution misses them, a real defect could ship undetected.

The defect itself is NOT injected automatically — each scenario documents what the defect
is and which code file to change.  The harness validates that the AI would have caught it
by checking its recommendations against the CR, not by running the application.

Usage (run from repo root)
--------------------------
    python3 harnesses/reliability/run.py
    python3 harnesses/reliability/run.py --config harnesses/reliability/config.json
    python3 harnesses/reliability/run.py --iterations 3
"""

from __future__ import annotations

import argparse
import asyncio
import dataclasses
import json
import logging
import os
import sys
import time
import uuid
from datetime import datetime, timezone
from typing import Optional

# ---------------------------------------------------------------------------
# Ensure repo root is on sys.path so `from shared.X import Y` works when this
# script is run as  python3 harnesses/reliability/run.py  from any cwd.
# ---------------------------------------------------------------------------
_HERE = os.path.dirname(os.path.abspath(__file__))
_REPO_ROOT = os.path.dirname(os.path.dirname(_HERE))
if _REPO_ROOT not in sys.path:
    sys.path.insert(0, _REPO_ROOT)

from shared.client import SolutionClient  # noqa: E402
from shared.dotenv import load_dotenv     # noqa: E402

logging.basicConfig(
    level=logging.INFO,
    format="%(asctime)s %(levelname)-8s %(name)s — %(message)s",
    datefmt="%H:%M:%S",
)
logger = logging.getLogger("reliability")


# ---------------------------------------------------------------------------
# Config
# ---------------------------------------------------------------------------

DEFAULT_CONFIG = {
    "agent_base_url": "http://localhost:8000",
    "integration_base_url": "http://localhost:3001",
    "tenant_id": "00000000-0000-0000-0000-000000000001",
    "output_dir": "harnesses/reliability/runs",
    "selector": {
        "source": "ado",
        "test_case_tags": ["Regression"],
        "integration_url": "",
        "integration_project": "",
        "integration_token": "",
    },
    "runner": {
        "iterations": 5,
        "tcg_enabled": True,
        "bypass_cache": True,
        "poll_interval_sec": 3.0,
        "poll_timeout_sec": 1200.0,
    },
    # response_schema lets harness adapt to the solution's field names
    "response_schema": {
        "action_field": "recommended_action",
        "risk_field": "overall_risk",
        "summary_key": "summary",
        "scored_tests_key": "scored_tests",
        "test_name_field": "test_name",
        "test_id_field": "test_id",
    },
    "scenarios": [],
}


def _deep_merge(base: dict, override: dict) -> None:
    for k, v in override.items():
        if isinstance(v, dict) and isinstance(base.get(k), dict):
            _deep_merge(base[k], v)
        else:
            base[k] = v


def load_config(path: Optional[str], iterations_override: Optional[int]) -> dict:
    cfg = json.loads(json.dumps(DEFAULT_CONFIG))
    if path and os.path.exists(path):
        with open(path, encoding="utf-8") as fh:
            _deep_merge(cfg, json.load(fh))

    load_dotenv()

    sel = cfg["selector"]
    sel["integration_url"] = os.getenv("ADO_BASE_URL", sel.get("integration_url", ""))
    sel["integration_project"] = os.getenv("ADO_PROJECT_KEY", sel.get("integration_project", ""))
    sel["integration_token"] = os.getenv("ADO_TOKEN", sel.get("integration_token", ""))

    if iterations_override is not None:
        cfg["runner"]["iterations"] = iterations_override

    return cfg


# ---------------------------------------------------------------------------
# Polling
# ---------------------------------------------------------------------------

async def _poll(client: SolutionClient, run_id: str, runner_cfg: dict) -> dict:
    deadline = time.monotonic() + runner_cfg.get("poll_timeout_sec", 1200.0)
    interval = runner_cfg.get("poll_interval_sec", 3.0)
    done = {"complete", "cache_hit"}
    while True:
        row = await client.poll(run_id)
        status = row.get("status", "unknown")
        if status in done:
            report = row.get("report")
            if report is None:
                raise RuntimeError(f"run {run_id} {status} but report is null")
            return report
        if status == "failed":
            raise RuntimeError(f"run {run_id} failed: {row.get('error_message')}")
        if time.monotonic() >= deadline:
            raise RuntimeError(f"run {run_id} still {status!r} after timeout")
        await asyncio.sleep(interval)


# ---------------------------------------------------------------------------
# Per-run result
# ---------------------------------------------------------------------------

@dataclasses.dataclass
class RunOutcome:
    index: int
    run_id: Optional[str]
    cr_id_used: str
    passed: bool
    violations: list[str]
    overall_risk: Optional[str]
    found_tests: list[str]       # tests flagged Must Run or Should Run
    missing_tests: list[str]     # expected tests that were NOT flagged
    elapsed_sec: float
    error: Optional[str] = None


def check_reliability(
    report: dict, expected_must_run: list[str], schema: dict
) -> tuple[list[str], list[str], list[str]]:
    """
    Returns (violations, found_tests, missing_tests).
    A test passes if it appears in scored_tests with Must Run or Should Run.
    """
    scored_key = schema.get("scored_tests_key", "scored_tests")
    action_field = schema.get("action_field", "recommended_action")
    name_field = schema.get("test_name_field", "test_name")
    id_field = schema.get("test_id_field", "test_id")

    scored = report.get(scored_key) or []
    flagged = {
        (t.get(name_field) or t.get(id_field) or "").strip()
        for t in scored
        if (t.get(action_field) or "").lower() in {"must run", "should run"}
    }

    found, missing, violations = [], [], []
    for expected in expected_must_run:
        # Partial match — test names may include suffixes
        matched = next((f for f in flagged if expected.lower() in f.lower() or f.lower() in expected.lower()), None)
        if matched:
            found.append(matched)
        else:
            missing.append(expected)
            violations.append(
                f"Expected '{expected}' to be Must Run or Should Run — "
                f"not found in flagged tests. Flagged: {sorted(flagged) or '(none)'}"
            )

    return violations, found, missing


async def run_once(
    client: SolutionClient,
    index: int,
    normalised_cr: dict,
    existing_tests: list,
    scenario: dict,
    cfg: dict,
) -> RunOutcome:
    runner_cfg = cfg["runner"]
    schema = cfg.get("response_schema", {})
    summary_key = schema.get("summary_key", "summary")
    risk_field = schema.get("risk_field", "overall_risk")

    cr_id = scenario["cr_id"]
    expected = scenario["expected_must_run"]
    bypass = runner_cfg.get("bypass_cache", True)
    cr_id_used = f"{cr_id}#{uuid.uuid4().hex[:8]}" if bypass else cr_id
    started = time.monotonic()

    try:
        submit = await client.submit(
            cr_id=cr_id_used,
            normalised_cr=normalised_cr,
            existing_tests=existing_tests,
            tcg_enabled=runner_cfg.get("tcg_enabled", True),
        )

        if submit.get("cache_hit"):
            report = {k: v for k, v in submit.items() if k != "cache_hit"}
            run_id = submit.get("run_id")
        else:
            run_id = submit.get("run_id")
            if not run_id:
                raise RuntimeError(f"no run_id in submit response: {submit!r}")
            logger.info("run %d submitted — run_id=%s", index + 1, run_id)
            report = await _poll(client, run_id, runner_cfg)

        violations, found, missing = check_reliability(report, expected, schema)
        summary = report.get(summary_key) or {}

        return RunOutcome(
            index=index,
            run_id=run_id,
            cr_id_used=cr_id_used,
            passed=len(violations) == 0,
            violations=violations,
            overall_risk=summary.get(risk_field),
            found_tests=found,
            missing_tests=missing,
            elapsed_sec=round(time.monotonic() - started, 2),
        )

    except Exception as exc:
        return RunOutcome(
            index=index,
            run_id=None,
            cr_id_used=cr_id_used,
            passed=False,
            violations=[f"Run error: {exc}"],
            overall_risk=None,
            found_tests=[],
            missing_tests=expected,
            elapsed_sec=round(time.monotonic() - started, 2),
            error=str(exc),
        )


# ---------------------------------------------------------------------------
# Report rendering
# ---------------------------------------------------------------------------

def _band(score: float) -> str:
    return "Excellent" if score == 1.0 else ("Good" if score >= 0.8 else ("Fair" if score >= 0.6 else "Poor"))


def render_scenario_section(scenario: dict, outcomes: list[RunOutcome]) -> list[str]:
    passed = [o for o in outcomes if o.passed]
    total = len(outcomes)
    score = len(passed) / total if total else 0.0

    lines = []
    lines.append(f"  {scenario['label']}")
    lines.append(f"  CR: {scenario['cr_id']}   Expected Must Run: {scenario['expected_must_run']}")
    lines.append(f"  Defect: {scenario.get('defect', 'N/A')}")
    lines.append(f"  {'─' * 56}")
    lines.append(f"  Reliability score: {score:.3f}  [{_band(score)}]   ({len(passed)}/{total} passed)")
    lines.append("")
    lines.append(f"  {'Run':<5} {'Status':<8} {'Risk':<10} {'Found':<30} {'Time':<8}")
    lines.append("  " + "-" * 60)
    for o in outcomes:
        status = "PASS" if o.passed else ("ERROR" if o.error else "FAIL")
        found = ", ".join(o.found_tests) or "(none)"
        lines.append(f"  {o.index+1:<5} {status:<8} {(o.overall_risk or '?'):<10} {found:<30} {o.elapsed_sec:.1f}s")

    bad = [o for o in outcomes if not o.passed]
    if bad:
        lines.append("")
        lines.append("  Violations")
        for o in bad:
            lines.append(f"  Run {o.index + 1}:")
            for v in o.violations:
                lines.append(f"    - {v}")

    return lines


def render_report(scenario_results: list[dict]) -> str:
    all_outcomes = [o for r in scenario_results for o in r["outcomes"]]
    total_all = len(all_outcomes)
    passed_all = sum(1 for o in all_outcomes if o.passed)
    overall = passed_all / total_all if total_all else 0.0

    lines = []
    lines.append("=" * 64)
    lines.append("  RELIABILITY HARNESS REPORT")
    lines.append("=" * 64)
    lines.append(f"  Scenarios:       {len(scenario_results)}")
    lines.append(f"  Total runs:      {total_all}")
    lines.append(f"  Overall passed:  {passed_all}")
    lines.append(f"  Overall score:   {overall:.3f}  [{_band(overall)}]")
    lines.append("")

    for r in scenario_results:
        lines.extend(render_scenario_section(r["scenario"], r["outcomes"]))
        lines.append("")

    lines.append("=" * 64)
    return "\n".join(lines)


def render_json(scenario_results: list[dict], cfg: dict) -> dict:
    all_outcomes = [o for r in scenario_results for o in r["outcomes"]]
    total_all = len(all_outcomes)
    passed_all = sum(1 for o in all_outcomes if o.passed)
    return {
        "timestamp": datetime.now(timezone.utc).isoformat(),
        "overall_reliability_score": round(passed_all / total_all, 4) if total_all else 0.0,
        "scenarios": [
            {
                "label": r["scenario"]["label"],
                "cr_id": r["scenario"]["cr_id"],
                "defect": r["scenario"].get("defect", ""),
                "expected_must_run": r["scenario"]["expected_must_run"],
                "iterations": len(r["outcomes"]),
                "passed": sum(1 for o in r["outcomes"] if o.passed),
                "reliability_score": round(
                    sum(1 for o in r["outcomes"] if o.passed) / len(r["outcomes"]), 4
                ) if r["outcomes"] else 0.0,
                "runs": [
                    {
                        "index": o.index,
                        "run_id": o.run_id,
                        "cr_id_used": o.cr_id_used,
                        "passed": o.passed,
                        "violations": o.violations,
                        "overall_risk": o.overall_risk,
                        "found_tests": o.found_tests,
                        "missing_tests": o.missing_tests,
                        "elapsed_sec": o.elapsed_sec,
                        "error": o.error,
                    }
                    for o in r["outcomes"]
                ],
            }
            for r in scenario_results
        ],
    }


# ---------------------------------------------------------------------------
# Word document report
# ---------------------------------------------------------------------------

def render_docx(scenario_results: list[dict], cfg: dict, out_path: str) -> None:
    try:
        from docx import Document as DocxDocument
        from docx.shared import Pt, RGBColor, Inches
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.oxml.ns import qn
        from docx.oxml import OxmlElement
    except ImportError:
        logger.warning("python-docx not installed — skipping Word report")
        return

    def rgb(hex_str: str) -> RGBColor:
        h = hex_str.lstrip("#")
        return RGBColor(int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16))

    NAVY  = "1A2F5A"
    BLUE  = "2471A3"
    GREEN = "1E8449"
    RED   = "C0392B"
    AMBER = "B7950B"
    WHITE = "FFFFFF"
    LGREY = "F2F3F4"

    def shade_cell(cell, hex_color: str) -> None:
        tc = cell._tc
        tcPr = tc.get_or_add_tcPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear")
        shd.set(qn("w:color"), "auto")
        shd.set(qn("w:fill"), hex_color)
        tcPr.append(shd)

    def set_col_widths(table, widths: list) -> None:
        for row in table.rows:
            for i, cell in enumerate(row.cells):
                cell.width = Inches(widths[i])

    def add_run(para, text, bold=False, color=None, size=None):
        r = para.add_run(text)
        r.bold = bold
        if color:
            r.font.color.rgb = rgb(color)
        if size:
            r.font.size = Pt(size)
        return r

    doc = DocxDocument()
    for section in doc.sections:
        section.top_margin = section.bottom_margin = Inches(0.8)
        section.left_margin = section.right_margin = Inches(1.0)

    all_outcomes = [o for r in scenario_results for o in r["outcomes"]]
    total_all = len(all_outcomes)
    passed_all = sum(1 for o in all_outcomes if o.passed)
    overall = passed_all / total_all if total_all else 0.0
    overall_band = _band(overall)
    overall_color = GREEN if overall >= 0.8 else (AMBER if overall >= 0.6 else RED)
    run_date = datetime.now(timezone.utc).strftime("%d %B %Y")
    llm_label = os.environ.get("LLM_PROVIDER", "").strip().upper() or "unknown"
    _model_map = {"OPENAI": "LLM_MODEL", "GEMINI": "GEMINI_MODEL", "OPENROUTER": "OPENROUTER_MODEL", "OLLAMA": "OLLAMA_MODEL"}
    model_label = os.environ.get(_model_map.get(llm_label, ""), "") or "unknown"

    # Title
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_run(title, "RELIABILITY HARNESS REPORT", bold=True, color=NAVY, size=22)

    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_run(sub, "AI Solution Testing Harness — Defect Detection Validation", color=BLUE, size=11)

    doc.add_paragraph()

    # Meta table
    meta = doc.add_table(rows=5, cols=2)
    for row, (label, value) in zip(meta.rows, [
        ("Scenarios tested", str(len(scenario_results))),
        ("Run date",         run_date),
        ("Total runs",       str(total_all)),
        ("LLM model",        f"{model_label} ({llm_label})"),
        ("Overall score",    f"{overall:.3f}  [{overall_band}]"),
    ]):
        shade_cell(row.cells[0], LGREY)
        lp = row.cells[0].paragraphs[0]
        lr = lp.add_run(label)
        lr.bold = True
        lr.font.color.rgb = rgb(NAVY)
        vp = row.cells[1].paragraphs[0]
        vr = vp.add_run(value)
        if label == "Overall score":
            vr.bold = True
            vr.font.color.rgb = rgb(overall_color)
    set_col_widths(meta, [1.8, 4.7])

    doc.add_paragraph()

    # Per-scenario sections
    for r in scenario_results:
        scenario = r["scenario"]
        outcomes = r["outcomes"]
        passed = [o for o in outcomes if o.passed]
        total = len(outcomes)
        score = len(passed) / total if total else 0.0
        score_color = GREEN if score >= 0.8 else (AMBER if score >= 0.6 else RED)

        h = doc.add_heading(scenario["label"], level=2)
        h.runs[0].font.color.rgb = rgb(BLUE)

        # Scenario detail
        detail = doc.add_table(rows=3, cols=2)
        for row2, (label, value) in zip(detail.rows, [
            ("CR",              scenario["cr_id"]),
            ("Defect",          scenario.get("defect", "")),
            ("Expected Must Run", ", ".join(scenario["expected_must_run"])),
        ]):
            shade_cell(row2.cells[0], LGREY)
            row2.cells[0].paragraphs[0].add_run(label).bold = True
            row2.cells[1].paragraphs[0].add_run(value)
        set_col_widths(detail, [1.8, 4.7])

        doc.add_paragraph()

        score_p = doc.add_paragraph()
        add_run(score_p, f"Reliability score: {score:.3f}  [{_band(score)}]  —  {len(passed)} of {total} runs passed",
                bold=True, color=GREEN if score >= 0.8 else RED)

        doc.add_paragraph()

        # Per-run table
        tbl = doc.add_table(rows=len(outcomes) + 1, cols=5)
        for cell, hdr in zip(tbl.rows[0].cells, ["Run", "Status", "Risk", "Tests Found as Must/Should Run", "Time"]):
            shade_cell(cell, NAVY)
            p2 = cell.paragraphs[0]
            p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r2 = p2.add_run(hdr)
            r2.bold = True
            r2.font.color.rgb = rgb(WHITE)

        for o, tbl_row in zip(outcomes, tbl.rows[1:]):
            status = "PASS" if o.passed else ("ERROR" if o.error else "FAIL")
            s_color = GREEN if o.passed else RED
            vals = [str(o.index+1), status, o.overall_risk or "?",
                    ", ".join(o.found_tests) or "(none)", f"{o.elapsed_sec:.1f}s"]
            colors = [None, s_color, None, GREEN if o.found_tests else RED, None]
            bolds = [False, True, False, False, False]
            for cell, val, col, bold in zip(tbl_row.cells, vals, colors, bolds):
                p3 = cell.paragraphs[0]
                p3.alignment = WD_ALIGN_PARAGRAPH.CENTER
                rn = p3.add_run(val)
                rn.bold = bold
                if col:
                    rn.font.color.rgb = rgb(col)
        set_col_widths(tbl, [0.5, 0.8, 1.0, 3.0, 0.8])

        failed_runs = [o for o in outcomes if not o.passed]
        if failed_runs:
            doc.add_paragraph()
            h3 = doc.add_heading("Violations", level=3)
            h3.runs[0].font.color.rgb = rgb(RED)
            for o in failed_runs:
                p4 = doc.add_paragraph()
                p4.add_run(f"Run {o.index + 1}:").bold = True
                for v in o.violations:
                    vp = doc.add_paragraph(style="List Bullet")
                    vp.add_run(v).font.color.rgb = rgb(RED)

        doc.add_paragraph()

    # Summary
    h_sum = doc.add_heading("Overall Summary", level=2)
    h_sum.runs[0].font.color.rgb = rgb(BLUE)
    verdict = doc.add_paragraph()
    add_run(verdict, f"Overall reliability score: {overall:.3f} [{overall_band}]  —  ", bold=True, color=overall_color)
    add_run(verdict, f"{passed_all} of {total_all} runs passed across {len(scenario_results)} scenarios.", size=12)

    if overall == 1.0:
        note = doc.add_paragraph()
        note.add_run("All scenarios passed. The solution correctly identified the Must Run tests on every run.").font.color.rgb = rgb(GREEN)
    else:
        note = doc.add_paragraph()
        note.add_run("One or more scenarios failed. The solution missed expected Must Run tests — review violations above.").font.color.rgb = rgb(RED)

    doc.save(out_path)
    logger.info("Word report written to %s", out_path)


# ---------------------------------------------------------------------------
# Main
# ---------------------------------------------------------------------------

async def main_async(cfg: dict) -> int:
    sel = cfg["selector"]
    runner_cfg = cfg["runner"]
    iterations = runner_cfg["iterations"]
    scenarios = cfg["scenarios"]
    endpoints = cfg.get("endpoints", {})

    if not scenarios:
        logger.error("No scenarios defined in config. Add scenarios to config.json.")
        return 1

    scenario_results: list[dict] = []

    async with SolutionClient(
        cfg["agent_base_url"],
        cfg["integration_base_url"],
        cfg["tenant_id"],
        endpoints=endpoints,
    ) as client:
        try:
            await client.health()
        except Exception as exc:
            logger.error("Agent service health check failed: %s", exc)
            logger.error("Is the stack running?")
            return 1

        # Fetch shared test cases once
        logger.info("Fetching existing test cases (tags: %s)...", sel.get("test_case_tags"))
        existing_tests = await client.fetch_test_cases_by_tags(
            tags=sel.get("test_case_tags") or [],
            url=sel.get("integration_url", ""),
            project=sel.get("integration_project", ""),
            token=sel.get("integration_token", ""),
        )
        logger.info("Loaded %d existing tests", len(existing_tests))

        for scenario in scenarios:
            cr_id = scenario["cr_id"]
            label = scenario["label"]
            outcomes: list[RunOutcome] = []

            logger.info("=" * 50)
            logger.info("Scenario: %s", label)

            # Use inline normalised_cr if defined — no integration fetch needed
            if "normalised_cr" in scenario:
                normalised_cr = scenario["normalised_cr"]
                logger.info("Using inline CR definition (no integration fetch)")
            else:
                logger.info("Fetching CR %s ...", cr_id)
                cr_data = await client.fetch_cr(
                    tool=sel.get("source", "ado"),
                    cr_id=cr_id,
                    url=sel.get("integration_url", ""),
                    project=sel.get("integration_project", ""),
                    token=sel.get("integration_token", ""),
                )
                normalised_cr = cr_data.get("normalised_cr") or {}

            cr_title = normalised_cr.get("title", "(no title)")
            logger.info("CR title: %s", cr_title)
            logger.info("Expected Must Run: %s", scenario["expected_must_run"])

            logger.info("Running %d times...", iterations)
            for i in range(iterations):
                logger.info("--- Run %d/%d ---", i + 1, iterations)
                outcome = await run_once(client, i, normalised_cr, existing_tests, scenario, cfg)
                status = "PASS" if outcome.passed else "FAIL"
                logger.info(
                    "Run %d: %s  risk=%s  found=%s  missing=%s  (%.1fs)",
                    i + 1, status,
                    outcome.overall_risk or "?",
                    outcome.found_tests or "(none)",
                    outcome.missing_tests or "(none)",
                    outcome.elapsed_sec,
                )
                if outcome.violations:
                    for v in outcome.violations:
                        logger.warning("  ! %s", v)
                outcomes.append(outcome)

            scenario_results.append({"scenario": scenario, "outcomes": outcomes})

    # Write artifacts
    out_dir = os.path.join(
        cfg["output_dir"],
        datetime.now(timezone.utc).strftime("%Y%m%dT%H%M%SZ"),
    )
    os.makedirs(out_dir, exist_ok=True)

    json_path = os.path.join(out_dir, "reliability_report.json")
    with open(json_path, "w", encoding="utf-8") as fh:
        json.dump(render_json(scenario_results, cfg), fh, indent=2)

    txt = render_report(scenario_results)
    txt_path = os.path.join(out_dir, "summary.txt")
    with open(txt_path, "w", encoding="utf-8") as fh:
        fh.write(txt)

    docx_path = os.path.join(out_dir, "reliability_report.docx")
    render_docx(scenario_results, cfg, docx_path)

    print()
    print(txt)
    print(f"\nArtifacts written to: {out_dir}/")
    print(f"  {json_path}")
    print(f"  {txt_path}")
    print(f"  {docx_path}")

    all_outcomes = [o for r in scenario_results for o in r["outcomes"]]
    passed_all = sum(1 for o in all_outcomes if o.passed)
    return 0 if passed_all == len(all_outcomes) else 1


def main() -> None:
    parser = argparse.ArgumentParser(description="Reliability harness for AI solution testing")
    parser.add_argument("--config", default="harnesses/reliability/config.json")
    parser.add_argument("--iterations", type=int, default=None)
    args = parser.parse_args()
    cfg = load_config(args.config, args.iterations)
    sys.exit(asyncio.run(main_async(cfg)))


if __name__ == "__main__":
    main()
