"""Integrity harness — asserts the AI solution does NOT hallucinate impact for
low-impact change requests.

Design intent
-------------
A CR that is genuinely low-impact should produce:
  • overall_risk = Low (or whichever values config allows)
  • no forbidden recommended_action values (e.g. Must Run / Should Run)
  • zero (or near-zero) generated test cases

If the solution repeatedly assigns High/Critical risk or flags Must Run tests for
such a CR, that is a hallucination of impact.  This harness quantifies that.

Usage (run from repo root)
--------------------------
    python3 harnesses/integrity/run.py --config harnesses/integrity/config.json
    python3 harnesses/integrity/run.py --iterations 10
"""

from __future__ import annotations

import argparse
import asyncio
import dataclasses
import json
import logging
import os
import sys
import time
import uuid
from datetime import datetime, timezone
from typing import Optional

# ---------------------------------------------------------------------------
# Ensure repo root is on sys.path so `from shared.X import Y` works when this
# script is run as  python3 harnesses/integrity/run.py  from any cwd.
# ---------------------------------------------------------------------------
_HERE = os.path.dirname(os.path.abspath(__file__))
_REPO_ROOT = os.path.dirname(os.path.dirname(_HERE))
if _REPO_ROOT not in sys.path:
    sys.path.insert(0, _REPO_ROOT)

from shared.client import SolutionClient  # noqa: E402
from shared.dotenv import load_dotenv     # noqa: E402

# ---------------------------------------------------------------------------
# Logging
# ---------------------------------------------------------------------------
logging.basicConfig(
    level=logging.INFO,
    format="%(asctime)s %(levelname)-8s %(name)s — %(message)s",
    datefmt="%H:%M:%S",
)
logger = logging.getLogger("integrity")


# ---------------------------------------------------------------------------
# Config
# ---------------------------------------------------------------------------

DEFAULT_CONFIG = {
    "agent_base_url": "http://localhost:8000",
    "integration_base_url": "http://localhost:3001",
    "tenant_id": "00000000-0000-0000-0000-000000000001",
    "output_dir": "harnesses/integrity/runs",
    "selector": {
        "mode": "ado",
        "cr_id": "158",
        "source": "ado",
        "test_case_tags": ["Regression"],
        "integration_url": "",
        "integration_project": "",
        "integration_token": "",
    },
    "runner": {
        "iterations": 5,
        "tcg_enabled": True,
        "bypass_cache": True,
        "poll_interval_sec": 3.0,
        "poll_timeout_sec": 1200.0,
    },
    # response_schema lets harness adapt to the solution's field names
    "response_schema": {
        "action_field": "recommended_action",
        "risk_field": "overall_risk",
        "summary_key": "summary",
        "scored_tests_key": "scored_tests",
        "generated_tcs_key": "generated_test_cases",
        "test_name_field": "test_name",
        "test_id_field": "test_id",
    },
    # Legacy single-CR format — used if "crs" array is absent
    "thresholds": {
        "allowed_overall_risk": ["Low"],
        "forbidden_actions": ["Must Run", "Should Run"],
        "max_generated_test_cases": 1,
    },
}


def _deep_merge(base: dict, override: dict) -> None:
    for k, v in override.items():
        if isinstance(v, dict) and isinstance(base.get(k), dict):
            _deep_merge(base[k], v)
        else:
            base[k] = v


def load_config(path: Optional[str], iterations_override: Optional[int]) -> dict:
    cfg = json.loads(json.dumps(DEFAULT_CONFIG))  # deep copy
    if path and os.path.exists(path):
        with open(path, encoding="utf-8") as fh:
            file_cfg = json.load(fh)
        _deep_merge(cfg, file_cfg)

    # Load .env so LLM vars are available when running outside Docker
    load_dotenv()

    # env overrides for integration credentials
    sel = cfg["selector"]
    sel["integration_url"] = os.getenv("ADO_BASE_URL", sel.get("integration_url", ""))
    sel["integration_project"] = os.getenv("ADO_PROJECT_KEY", sel.get("integration_project", ""))
    sel["integration_token"] = os.getenv("ADO_TOKEN", sel.get("integration_token", ""))

    if iterations_override is not None:
        cfg["runner"]["iterations"] = iterations_override

    # Normalise legacy single-CR format into a "crs" list so the rest of the
    # code only needs to handle one shape.
    if "crs" not in cfg:
        cr_id = cfg.get("selector", {}).get("cr_id") or cfg.get("cr_id", "unknown")
        cfg["crs"] = [{"cr_id": cr_id, "thresholds": cfg["thresholds"]}]

    return cfg


# ---------------------------------------------------------------------------
# Per-run integrity check
# ---------------------------------------------------------------------------

RISK_RANK = {"Low": 1, "Medium": 2, "High": 3, "Critical": 4}


@dataclasses.dataclass
class RunOutcome:
    index: int
    run_id: Optional[str]
    cr_id_used: str
    passed: bool
    violations: list[str]
    overall_risk: Optional[str]
    must_or_should_count: int
    generated_tc_count: int
    elapsed_sec: float
    error: Optional[str] = None


def check_integrity(report: dict, thresholds: dict, schema: dict) -> list[str]:
    """Return list of violation strings (empty = pass)."""
    violations: list[str] = []

    summary_key = schema.get("summary_key", "summary")
    risk_field = schema.get("risk_field", "overall_risk")
    action_field = schema.get("action_field", "recommended_action")
    scored_key = schema.get("scored_tests_key", "scored_tests")
    gen_tcs_key = schema.get("generated_tcs_key", "generated_test_cases")
    name_field = schema.get("test_name_field", "test_name")
    id_field = schema.get("test_id_field", "test_id")

    summary = report.get(summary_key) or {}
    overall_risk = summary.get(risk_field) or "Unknown"
    allowed = thresholds.get("allowed_overall_risk", ["Low"])
    if overall_risk not in allowed:
        violations.append(
            f"{risk_field}={overall_risk!r} — expected one of {allowed}"
        )

    forbidden_actions = set(thresholds.get("forbidden_actions", ["Must Run", "Should Run"]))
    scored = report.get(scored_key) or []
    bad_tests = [
        t for t in scored
        if (t.get(action_field) or "") in forbidden_actions
    ]
    if bad_tests:
        names = [t.get(name_field) or t.get(id_field) or "?" for t in bad_tests]
        violations.append(
            f"{len(bad_tests)} test(s) have forbidden {action_field} "
            f"({', '.join(sorted(forbidden_actions))}): {names[:5]}"
        )

    gen_tc = report.get(gen_tcs_key) or []
    max_tc = thresholds.get("max_generated_test_cases", 1)
    if len(gen_tc) > max_tc:
        violations.append(
            f"{len(gen_tc)} generated test cases — expected <= {max_tc} for a low-impact CR"
        )

    return violations


# ---------------------------------------------------------------------------
# Core async runner
# ---------------------------------------------------------------------------

async def run_once(
    client: SolutionClient,
    index: int,
    normalised_cr: dict,
    existing_tests: list,
    cfg: dict,
) -> RunOutcome:
    runner_cfg = cfg["runner"]
    thresholds = cfg["thresholds"]
    schema = cfg.get("response_schema", {})
    base_cr_id = cfg["selector"]["cr_id"]

    summary_key = schema.get("summary_key", "summary")
    risk_field = schema.get("risk_field", "overall_risk")
    action_field = schema.get("action_field", "recommended_action")
    scored_key = schema.get("scored_tests_key", "scored_tests")
    gen_tcs_key = schema.get("generated_tcs_key", "generated_test_cases")

    bypass = runner_cfg.get("bypass_cache", True)
    cr_id = f"{base_cr_id}#{uuid.uuid4().hex[:8]}" if bypass else base_cr_id
    started = time.monotonic()

    try:
        submit = await client.submit(
            cr_id=cr_id,
            normalised_cr=normalised_cr,
            existing_tests=existing_tests,
            tcg_enabled=runner_cfg.get("tcg_enabled", True),
        )

        if submit.get("cache_hit"):
            report = {k: v for k, v in submit.items() if k != "cache_hit"}
            run_id = submit.get("run_id")
            logger.warning("run %d hit cache (cr_id=%s)", index + 1, cr_id)
        else:
            run_id = submit.get("run_id")
            if not run_id:
                raise RuntimeError(f"no run_id in submit response: {submit!r}")
            logger.info("run %d submitted — run_id=%s", index + 1, run_id)
            report = await _poll(client, run_id, runner_cfg)

        violations = check_integrity(report, thresholds, schema)
        summary = report.get(summary_key) or {}
        scored = report.get(scored_key) or []
        forbidden_actions = set(thresholds.get("forbidden_actions", ["Must Run", "Should Run"]))
        must_should = sum(
            1 for t in scored
            if (t.get(action_field) or "") in forbidden_actions
        )
        gen_tc = len(report.get(gen_tcs_key) or [])

        return RunOutcome(
            index=index,
            run_id=run_id,
            cr_id_used=cr_id,
            passed=len(violations) == 0,
            violations=violations,
            overall_risk=summary.get(risk_field),
            must_or_should_count=must_should,
            generated_tc_count=gen_tc,
            elapsed_sec=round(time.monotonic() - started, 2),
        )

    except Exception as exc:  # noqa: BLE001
        return RunOutcome(
            index=index,
            run_id=None,
            cr_id_used=cr_id,
            passed=False,
            violations=[f"Run error: {exc}"],
            overall_risk=None,
            must_or_should_count=0,
            generated_tc_count=0,
            elapsed_sec=round(time.monotonic() - started, 2),
            error=str(exc),
        )


async def _poll(client: SolutionClient, run_id: str, runner_cfg: dict) -> dict:
    deadline = time.monotonic() + runner_cfg.get("poll_timeout_sec", 1200.0)
    interval = runner_cfg.get("poll_interval_sec", 3.0)
    done_statuses = {"complete", "cache_hit"}
    failed_statuses = {"failed"}

    while True:
        row = await client.poll(run_id)
        status = row.get("status", "unknown")
        if status in done_statuses:
            report = row.get("report")
            if report is None:
                raise RuntimeError(f"run {run_id} {status} but report is null")
            return report
        if status in failed_statuses:
            raise RuntimeError(f"run {run_id} failed: {row.get('error_message')}")
        if time.monotonic() >= deadline:
            raise RuntimeError(f"run {run_id} still {status!r} after timeout")
        await asyncio.sleep(interval)


# ---------------------------------------------------------------------------
# Report rendering
# ---------------------------------------------------------------------------

def _score_band(score: float) -> str:
    return "Excellent" if score == 1.0 else ("Good" if score >= 0.8 else ("Fair" if score >= 0.6 else "Poor"))


def render_cr_section(cr_id: str, cr_title: str, outcomes: list[RunOutcome], thresholds: dict) -> list[str]:
    """Return lines for a single CR block inside the combined report."""
    passed = [o for o in outcomes if o.passed]
    failed = [o for o in outcomes if not o.passed and not o.error]
    errored = [o for o in outcomes if o.error]
    total = len(outcomes)
    score = len(passed) / total if total else 0.0
    band = _score_band(score)

    lines: list[str] = []
    lines.append(f"  CR {cr_id}  —  {cr_title}")
    lines.append(f"  {'─' * 56}")
    lines.append(f"  Runs: {total}   Passed: {len(passed)}   Failed: {len(failed)}   Errored: {len(errored)}")
    lines.append(f"  Integrity score: {score:.3f}  [{band}]")
    lines.append(f"  Thresholds: risk={thresholds['allowed_overall_risk']}  "
                 f"forbidden={thresholds['forbidden_actions']}  "
                 f"max_tcs={thresholds['max_generated_test_cases']}")
    lines.append("")
    lines.append(f"  {'Run':<5} {'Status':<8} {'Risk':<10} {'Must/Should':<13} {'Gen TCs':<9} {'Time':<8}")
    lines.append("  " + "-" * 56)
    for o in outcomes:
        status = "PASS" if o.passed else ("ERROR" if o.error else "FAIL")
        lines.append(
            f"  {o.index + 1:<5} {status:<8} {(o.overall_risk or '?'):<10} "
            f"{o.must_or_should_count:<13} {o.generated_tc_count:<9} {o.elapsed_sec:.1f}s"
        )

    bad = [o for o in outcomes if not o.passed]
    if bad:
        lines.append("")
        lines.append("  Violations")
        for o in bad:
            lines.append(f"  Run {o.index + 1}:")
            for v in o.violations:
                lines.append(f"    - {v}")

    return lines


def render_report(cr_results: list[dict]) -> str:
    all_outcomes = [o for r in cr_results for o in r["outcomes"]]
    total_all = len(all_outcomes)
    passed_all = sum(1 for o in all_outcomes if o.passed)
    overall_score = passed_all / total_all if total_all else 0.0
    overall_band = _score_band(overall_score)

    lines: list[str] = []
    lines.append("=" * 60)
    lines.append("  INTEGRITY HARNESS REPORT")
    lines.append("=" * 60)
    lines.append(f"  CRs tested:      {len(cr_results)}")
    lines.append(f"  Total runs:      {total_all}")
    lines.append(f"  Overall passed:  {passed_all}")
    lines.append(f"  Overall score:   {overall_score:.3f}  [{overall_band}]")
    lines.append("")

    for r in cr_results:
        lines.extend(render_cr_section(r["cr_id"], r["cr_title"], r["outcomes"], r["thresholds"]))
        lines.append("")

    lines.append("=" * 60)
    return "\n".join(lines)


def render_json(cr_results: list[dict], cfg: dict) -> dict:
    all_outcomes = [o for r in cr_results for o in r["outcomes"]]
    total_all = len(all_outcomes)
    passed_all = sum(1 for o in all_outcomes if o.passed)
    return {
        "timestamp": datetime.now(timezone.utc).isoformat(),
        "overall_integrity_score": round(passed_all / total_all, 4) if total_all else 0.0,
        "crs": [
            {
                "cr_id": r["cr_id"],
                "cr_title": r["cr_title"],
                "iterations": len(r["outcomes"]),
                "passed": sum(1 for o in r["outcomes"] if o.passed),
                "failed": sum(1 for o in r["outcomes"] if not o.passed),
                "integrity_score": round(
                    sum(1 for o in r["outcomes"] if o.passed) / len(r["outcomes"]), 4
                ) if r["outcomes"] else 0.0,
                "thresholds": r["thresholds"],
                "runs": [
                    {
                        "index": o.index,
                        "run_id": o.run_id,
                        "cr_id_used": o.cr_id_used,
                        "passed": o.passed,
                        "violations": o.violations,
                        "overall_risk": o.overall_risk,
                        "must_or_should_count": o.must_or_should_count,
                        "generated_tc_count": o.generated_tc_count,
                        "elapsed_sec": o.elapsed_sec,
                        "error": o.error,
                    }
                    for o in r["outcomes"]
                ],
            }
            for r in cr_results
        ],
    }


# ---------------------------------------------------------------------------
# Word document report
# ---------------------------------------------------------------------------

def render_docx(cr_results: list[dict], cfg: dict, out_path: str) -> None:
    """cr_results: list of {"cr_id", "cr_title", "outcomes", "thresholds"}"""
    try:
        from docx import Document as DocxDocument
        from docx.shared import Pt, RGBColor, Inches
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.oxml.ns import qn
        from docx.oxml import OxmlElement
    except ImportError:
        logger.warning("python-docx not installed — skipping Word report. Run: pip3 install python-docx")
        return

    def rgb(hex_str: str) -> RGBColor:
        h = hex_str.lstrip("#")
        return RGBColor(int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16))

    NAVY   = "1A2F5A"
    BLUE   = "2471A3"
    GREEN  = "1E8449"
    RED    = "C0392B"
    AMBER  = "B7950B"
    WHITE  = "FFFFFF"
    LGREY  = "F2F3F4"

    def shade_cell(cell, hex_color: str) -> None:
        tc = cell._tc
        tcPr = tc.get_or_add_tcPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear")
        shd.set(qn("w:color"), "auto")
        shd.set(qn("w:fill"), hex_color)
        tcPr.append(shd)

    def set_col_widths(table, widths_inches: list) -> None:
        for row in table.rows:
            for i, cell in enumerate(row.cells):
                cell.width = Inches(widths_inches[i])

    doc = DocxDocument()

    # Page margins
    for section in doc.sections:
        section.top_margin    = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin   = Inches(1.0)
        section.right_margin  = Inches(1.0)

    all_outcomes = [o for r in cr_results for o in r["outcomes"]]
    total_all   = len(all_outcomes)
    passed_all  = sum(1 for o in all_outcomes if o.passed)
    overall_score = passed_all / total_all if total_all else 0.0
    overall_band  = _score_band(overall_score)
    overall_color = GREEN if overall_score >= 0.8 else (AMBER if overall_score >= 0.6 else RED)
    run_date = datetime.now(timezone.utc).strftime("%d %B %Y")
    llm_label = os.environ.get("LLM_PROVIDER", "").strip().lower()
    _provider_model_map = {
        "openai":      os.environ.get("LLM_MODEL", ""),
        "gemini":      os.environ.get("GEMINI_MODEL", ""),
        "openrouter":  os.environ.get("OPENROUTER_MODEL", ""),
        "ollama":      os.environ.get("OLLAMA_MODEL", ""),
    }
    model_label = _provider_model_map.get(llm_label, "") or "unknown"
    llm_label = llm_label.upper() or "unknown"

    # Title block
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("INTEGRITY HARNESS REPORT")
    run.bold = True
    run.font.size = Pt(22)
    run.font.color.rgb = rgb(NAVY)

    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sr = sub.add_run("AI Solution Testing Harness — Multi-CR Integrity Validation")
    sr.font.size = Pt(11)
    sr.font.color.rgb = rgb(BLUE)

    doc.add_paragraph()

    # Meta table
    meta = doc.add_table(rows=5, cols=2)
    meta.style = "Table Grid"
    cr_list = ", ".join(f"CR-{r['cr_id']}" for r in cr_results)
    for row, (label, value) in zip(meta.rows, [
        ("CRs tested",      cr_list),
        ("Run date",        run_date),
        ("Total runs",      str(total_all)),
        ("LLM model",       f"{model_label} ({llm_label})"),
        ("Overall score",   f"{overall_score:.3f}  [{overall_band}]"),
    ]):
        shade_cell(row.cells[0], LGREY)
        lp = row.cells[0].paragraphs[0]
        lr = lp.add_run(label)
        lr.bold = True
        lr.font.color.rgb = rgb(NAVY)
        vp = row.cells[1].paragraphs[0]
        vr = vp.add_run(value)
        if label == "Overall score":
            vr.bold = True
            vr.font.color.rgb = rgb(overall_color)
    set_col_widths(meta, [1.8, 4.7])

    doc.add_paragraph()

    # Per-CR sections
    for r in cr_results:
        outcomes = r["outcomes"]
        cr_id    = r["cr_id"]
        cr_title = r["cr_title"]
        t        = r["thresholds"]
        passed   = [o for o in outcomes if o.passed]
        total    = len(outcomes)
        score    = len(passed) / total if total else 0.0
        band     = _score_band(score)
        score_color = GREEN if score >= 0.8 else (AMBER if score >= 0.6 else RED)

        h_cr = doc.add_heading(f"CR {cr_id}  —  {cr_title}", level=2)
        h_cr.runs[0].font.color.rgb = rgb(BLUE)

        thr = doc.add_table(rows=4, cols=2)
        thr.style = "Table Grid"
        for cell, text in zip(thr.rows[0].cells, ["Threshold", "Pass condition"]):
            shade_cell(cell, NAVY)
            p = cell.paragraphs[0]
            r2 = p.add_run(text)
            r2.bold = True
            r2.font.color.rgb = rgb(WHITE)
        for thr_row, (label, value) in zip(thr.rows[1:], [
            ("overall_risk",         f"Must be one of: {t['allowed_overall_risk']}"),
            ("recommended_action",   f"No {', '.join(t['forbidden_actions'])}"),
            ("generated_test_cases", f"<= {t['max_generated_test_cases']}"),
        ]):
            thr_row.cells[0].paragraphs[0].add_run(label).bold = True
            thr_row.cells[1].paragraphs[0].add_run(value)
        set_col_widths(thr, [2.5, 4.0])

        doc.add_paragraph()

        score_p = doc.add_paragraph()
        score_p.add_run("Integrity score: ").bold = True
        sc_run = score_p.add_run(f"{score:.3f}  [{band}]  —  {len(passed)} of {total} runs passed")
        sc_run.bold = True
        sc_run.font.color.rgb = rgb(score_color)

        doc.add_paragraph()

        results_tbl = doc.add_table(rows=len(outcomes) + 1, cols=6)
        results_tbl.style = "Table Grid"
        headers = ["Run", "Status", "Overall Risk", "Must/Should", "Gen TCs", "Time"]
        for cell, hdr in zip(results_tbl.rows[0].cells, headers):
            shade_cell(cell, NAVY)
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run2 = p.add_run(hdr)
            run2.bold = True
            run2.font.color.rgb = rgb(WHITE)

        for o, tbl_row in zip(outcomes, results_tbl.rows[1:]):
            status   = "PASS" if o.passed else ("ERROR" if o.error else "FAIL")
            s_color  = GREEN if o.passed else RED
            r_color  = GREEN if o.overall_risk == "Low" else (AMBER if o.overall_risk == "Medium" else RED)
            values   = [
                str(o.index + 1), status, o.overall_risk or "?",
                str(o.must_or_should_count), str(o.generated_tc_count), f"{o.elapsed_sec:.1f}s",
            ]
            colors = [None, s_color, r_color, None, None, None]
            bolds  = [False, True, True, False, False, False]
            for cell, val, col, bold in zip(tbl_row.cells, values, colors, bolds):
                p = cell.paragraphs[0]
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                rn = p.add_run(val)
                rn.bold = bold
                if col:
                    rn.font.color.rgb = rgb(col)
        set_col_widths(results_tbl, [0.5, 0.9, 1.2, 1.2, 1.0, 0.8])

        doc.add_paragraph()

        failed_runs = [o for o in outcomes if not o.passed]
        if failed_runs:
            h_v = doc.add_heading("Violations", level=3)
            h_v.runs[0].font.color.rgb = rgb(RED)
            for o in failed_runs:
                p = doc.add_paragraph()
                rn = p.add_run(f"Run {o.index + 1}:")
                rn.bold = True
                for v in o.violations:
                    vp = doc.add_paragraph(style="List Bullet")
                    vr = vp.add_run(v)
                    vr.font.color.rgb = rgb(RED)
            doc.add_paragraph()

    # Overall summary
    h_sum = doc.add_heading("Overall Summary", level=2)
    h_sum.runs[0].font.color.rgb = rgb(BLUE)

    verdict = doc.add_paragraph()
    v1 = verdict.add_run(f"Overall integrity score: {overall_score:.3f} [{overall_band}]  —  ")
    v1.bold = True
    v1.font.color.rgb = rgb(overall_color)
    v2 = verdict.add_run(f"{passed_all} of {total_all} runs passed across {len(cr_results)} CRs.")
    v2.font.size = Pt(12)

    if overall_score == 1.0:
        note = doc.add_paragraph()
        note.add_run("All runs passed across all CRs. The solution performed correctly on every run.").font.color.rgb = rgb(GREEN)
    else:
        note = doc.add_paragraph()
        nr = note.add_run("One or more runs failed. Review the violations above.")
        nr.font.color.rgb = rgb(RED)

    doc.save(out_path)
    logger.info("Word report written to %s", out_path)


# ---------------------------------------------------------------------------
# Main
# ---------------------------------------------------------------------------

async def main_async(cfg: dict) -> int:
    sel = cfg["selector"]
    runner_cfg = cfg["runner"]
    iterations = runner_cfg["iterations"]
    cr_configs = cfg["crs"]  # list of {"cr_id": ..., "thresholds": ...}
    cr_results: list[dict] = []

    endpoints = cfg.get("endpoints", {})

    async with SolutionClient(
        cfg["agent_base_url"],
        cfg["integration_base_url"],
        cfg["tenant_id"],
        endpoints=endpoints,
    ) as client:
        # Health check
        try:
            await client.health()
        except Exception as exc:
            logger.error("Agent service health check failed: %s", exc)
            logger.error("Is the stack running?")
            return 1

        # Fetch shared test cases once
        logger.info("Fetching existing test cases (tags: %s)...", sel.get("test_case_tags"))
        existing_tests = await client.fetch_test_cases_by_tags(
            tags=sel.get("test_case_tags") or [],
            url=sel.get("integration_url", ""),
            project=sel.get("integration_project", ""),
            token=sel.get("integration_token", ""),
        )
        logger.info("Loaded %d existing tests", len(existing_tests))

        # Run each CR
        for cr_cfg in cr_configs:
            cr_id = cr_cfg["cr_id"]
            thresholds = cr_cfg["thresholds"]
            outcomes: list[RunOutcome] = []

            logger.info("=" * 50)
            logger.info("Fetching CR %s ...", cr_id)
            cr_data = await client.fetch_cr(
                tool=sel.get("source", "ado"),
                cr_id=cr_id,
                url=sel.get("integration_url", ""),
                project=sel.get("integration_project", ""),
                token=sel.get("integration_token", ""),
            )
            normalised_cr = cr_data.get("normalised_cr") or {}
            cr_title = normalised_cr.get("title", "(no title)")
            logger.info("CR title: %s", cr_title)

            # Build a per-CR config slice for run_once
            cr_run_cfg = {
                **cfg,
                "selector": {**sel, "cr_id": cr_id},
                "thresholds": thresholds,
            }

            logger.info("Running integrity check %d times for CR %s...", iterations, cr_id)
            for i in range(iterations):
                logger.info("--- CR %s  Run %d/%d ---", cr_id, i + 1, iterations)
                outcome = await run_once(client, i, normalised_cr, existing_tests, cr_run_cfg)
                status = "PASS" if outcome.passed else "FAIL"
                logger.info(
                    "Run %d: %s  risk=%s  must/should=%d  gen_tcs=%d  (%.1fs)",
                    i + 1, status,
                    outcome.overall_risk or "?",
                    outcome.must_or_should_count,
                    outcome.generated_tc_count,
                    outcome.elapsed_sec,
                )
                if outcome.violations:
                    for v in outcome.violations:
                        logger.warning("  ! %s", v)
                outcomes.append(outcome)

            cr_results.append({
                "cr_id": cr_id,
                "cr_title": cr_title,
                "outcomes": outcomes,
                "thresholds": thresholds,
            })

    # Write artifacts
    out_dir = os.path.join(
        cfg["output_dir"],
        datetime.now(timezone.utc).strftime("%Y%m%dT%H%M%SZ"),
    )
    os.makedirs(out_dir, exist_ok=True)

    report_json = render_json(cr_results, cfg)
    json_path = os.path.join(out_dir, "integrity_report.json")
    with open(json_path, "w", encoding="utf-8") as fh:
        json.dump(report_json, fh, indent=2)

    txt = render_report(cr_results)
    txt_path = os.path.join(out_dir, "summary.txt")
    with open(txt_path, "w", encoding="utf-8") as fh:
        fh.write(txt)

    docx_path = os.path.join(out_dir, "integrity_report.docx")
    render_docx(cr_results, cfg, docx_path)

    print()
    print(txt)
    print(f"\nArtifacts written to: {out_dir}/")
    print(f"  {json_path}")
    print(f"  {txt_path}")
    print(f"  {docx_path}")

    all_outcomes = [o for r in cr_results for o in r["outcomes"]]
    passed_all = sum(1 for o in all_outcomes if o.passed)
    return 0 if passed_all == len(all_outcomes) else 1


def main() -> None:
    parser = argparse.ArgumentParser(description="Integrity harness for AI solution testing")
    parser.add_argument("--config", default="harnesses/integrity/config.json", help="Path to config JSON")
    parser.add_argument("--iterations", type=int, default=None, help="Override number of runs")
    args = parser.parse_args()

    cfg = load_config(args.config, args.iterations)
    sys.exit(asyncio.run(main_async(cfg)))


if __name__ == "__main__":
    main()
