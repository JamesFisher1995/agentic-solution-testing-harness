"""Render a human-readable Markdown digest and Word report of a consistency run."""

from __future__ import annotations

import logging
from datetime import datetime, timezone
from typing import TYPE_CHECKING

from .narrative import generate_narrative, generate_improvements, generate_tc_topic_similarity

if TYPE_CHECKING:
    from .agents.base import SelectedCR
    from .agents.evaluator import ConsistencyReport
    from .agents.runner import RunResult
    from .config import Config

logger = logging.getLogger("harness.report")

_DIM_LABELS = {
    "test_set_stability":    ("Same tests selected",       "Are the same tests chosen every run?"),
    "action_agreement":      ("Same recommended action",   "Does each test get the same Must/Should/Consider/Exclude every run?"),
    "score_stability":       ("Stable relevance scores",   "Is the LLM's underlying relevance score for each test consistent? Score drift near a threshold can cause the action to flip between Must/Should/Consider."),
    "ranking_stability":     ("Stable ranking",            "Does the overall order of all tests stay the same across runs?"),
    "risk_agreement":        ("Stable overall risk",       "Is the overall risk rating (Critical/High/Medium/Low) the same every run?"),
    "summary_stability":     ("Stable headline counts",    "Are the must-run, should-run and exclude counts consistent across runs?"),
    "coverage_gap_stability":        ("Same coverage gaps",             "Are the same gaps in test coverage identified every run?"),
    "generated_test_case_consistency":("Consistent generated test cases", "Are a consistent number of new test cases generated, covering the same topics every run?"),
    "entity_stability":              ("Stable entity extraction",         "Does the story parser extract the same components/systems from the CR every run?"),
}


def render_summary_md(
    report: "ConsistencyReport",
    selected: "SelectedCR",
    results: "list[RunResult]",
    cfg: "Config",
) -> str:
    cr_title = (selected.normalised_cr or {}).get("title", "")
    run_date = datetime.now(timezone.utc).strftime("%Y-%m-%d")

    lines: list[str] = []
    lines.append(f"# Consistency report — CR {selected.cr_id}\n")

    # At-a-glance summary block
    lines.append(f"**CR:** {selected.cr_id}" + (f" — {cr_title}" if cr_title else ""))
    lines.append(
        f"**Ran:** {cfg.runner.iterations} iterations · "
        f"{report.runs_evaluated} completed · "
        f"{report.runs_failed} failed"
    )
    lines.append(f"**Date:** {run_date}")
    lines.append(f"**Source:** {selected.source.upper()} · "
                 f"**Tests scored:** {len(selected.existing_tests)} · "
                 f"**Cache bypass:** {'on' if cfg.runner.bypass_cache else 'off'}")
    lines.append("")

    if not report.sufficient:
        lines.append("> **Insufficient data** — need at least 2 successful runs.\n")
        _append_notes(lines, report)
        return "\n".join(lines)

    lines.append(f"## Overall: {report.overall_score:.3f} — {report.band}\n")
    lines.append(_bar(report.overall_score) + "\n")

    lines.append("## Dimensions\n")
    lines.append("| Dimension | What it measures | Score | Weight |")
    lines.append("|---|---|---|---|")
    sorted_dims = sorted(
        report.dimensions.items(),
        key=lambda kv: -(report.weights_used.get(kv[0]) or 0),
    )
    for key, val in sorted_dims:
        label, description = _DIM_LABELS.get(key, (key, ""))
        weight = report.weights_used.get(key)
        val_s = "—" if val is None else f"{val:.3f}"
        w_s = "excluded" if weight is None else f"{weight:.2f}"
        lines.append(f"| {label} | {description} | {val_s} | {w_s} |")
    lines.append("")

    d = report.details
    lines.append("## Test selection\n")
    lines.append(f"- Unique tests selected across all runs: **{d.get('union_test_count', 0)}**")
    lines.append(f"- Tests selected by *every* run: **{d.get('common_test_count', 0)}**")

    lines.append("")


    lines.append("## Classification & summary\n")
    lines.append(f"- overall_risk per run: `{d.get('overall_risk_values')}`")

    tc_counts = d.get("generated_tc_counts")
    if tc_counts is not None:
        lines.append(f"- Generated test cases per run: `{tc_counts}`")
        topic_verdict = generate_tc_topic_similarity(report)
        if topic_verdict:
            lines.append(f"- Topic consistency: {topic_verdict}")
    lines.append("")

    narrative = generate_narrative(report, selected, results)
    if narrative:
        lines.append(narrative)

    improvements = generate_improvements(report, selected)
    if improvements:
        lines.append(improvements)

    _append_notes(lines, report)
    return "\n".join(lines)


def _append_notes(lines: list[str], report: "ConsistencyReport") -> None:
    if report.notes:
        lines.append("## Notes\n")
        for n in report.notes:
            lines.append(f"- {n}")
        lines.append("")


def _append_run_table(lines: list[str], results: "list[RunResult]") -> None:
    lines.append("## Runs\n")
    lines.append("| # | cr_id used | status | run_id | elapsed (s) | error |")
    lines.append("|---|---|---|---|---|---|")
    for r in results:
        lines.append(
            f"| {r.index + 1} | {r.cr_id_used} | {r.status} | {r.run_id or '—'} | "
            f"{r.elapsed_sec} | {r.error or ''} |"
        )
    lines.append("")


def _bar(score: float, width: int = 30) -> str:
    filled = round(score * width)
    return "`[" + "#" * filled + "-" * (width - filled) + f"]` {score * 100:.1f}%"


# ---------------------------------------------------------------------------
# Word document report
# ---------------------------------------------------------------------------

def render_docx(
    report: "ConsistencyReport",
    selected: "SelectedCR",
    results: "list[RunResult]",
    cfg: "Config",
    out_path: str,
) -> None:
    try:
        from docx import Document as DocxDocument
        from docx.shared import Pt, RGBColor, Inches
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.oxml.ns import qn
        from docx.oxml import OxmlElement
    except ImportError:
        logger.warning("python-docx not installed — skipping Word report. Run: pip3 install python-docx")
        return

    def rgb(hex_str: str) -> RGBColor:
        h = hex_str.lstrip("#")
        return RGBColor(int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16))

    NAVY  = "1A2F5A"
    BLUE  = "2471A3"
    GREEN = "1E8449"
    RED   = "C0392B"
    AMBER = "B7950B"
    WHITE = "FFFFFF"
    LGREY = "F2F3F4"

    def shade_cell(cell, hex_color: str) -> None:
        tc = cell._tc
        tcPr = tc.get_or_add_tcPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear")
        shd.set(qn("w:color"), "auto")
        shd.set(qn("w:fill"), hex_color)
        tcPr.append(shd)

    def set_col_widths(table, widths_inches: list) -> None:
        for row in table.rows:
            for i, cell in enumerate(row.cells):
                cell.width = Inches(widths_inches[i])

    cr_title = (selected.normalised_cr or {}).get("title", "")
    run_date = datetime.now(timezone.utc).strftime("%d %B %Y")
    score = report.overall_score or 0.0
    band = report.band or "—"
    score_color = GREEN if score >= 0.90 else (AMBER if score >= 0.75 else RED)

    doc = DocxDocument()
    for section in doc.sections:
        section.top_margin    = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin   = Inches(1.0)
        section.right_margin  = Inches(1.0)

    # ── Title block ──────────────────────────────────────────────────────────
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    tr = title.add_run("CONSISTENCY HARNESS REPORT")
    tr.bold = True
    tr.font.size = Pt(22)
    tr.font.color.rgb = rgb(NAVY)

    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sr = sub.add_run("AI Test Impact Analyser — Output Consistency Evaluation")
    sr.font.size = Pt(11)
    sr.font.color.rgb = rgb(BLUE)

    doc.add_paragraph()

    # ── Meta table ───────────────────────────────────────────────────────────
    meta = doc.add_table(rows=6, cols=2)
    meta.style = "Table Grid"
    completed = [r for r in results if r.status in ("complete", "cache_hit")]
    for row, (label, value) in zip(meta.rows, [
        ("CR",                  f"{selected.cr_id}" + (f" — {cr_title}" if cr_title else "")),
        ("Run date",            run_date),
        ("Iterations",          f"{cfg.runner.iterations} configured · {report.runs_evaluated} completed · {report.runs_failed} failed"),
        ("Tests scored",        str(len(selected.existing_tests))),
        ("LLM model",           "gpt-4o-mini (OpenAI)"),
        ("Consistency score",   f"{score:.3f}  [{band}]"),
    ]):
        shade_cell(row.cells[0], LGREY)
        lp = row.cells[0].paragraphs[0]
        lr = lp.add_run(label)
        lr.bold = True
        lr.font.color.rgb = rgb(NAVY)
        vp = row.cells[1].paragraphs[0]
        vr = vp.add_run(value)
        if label == "Consistency score":
            vr.bold = True
            vr.font.color.rgb = rgb(score_color)
    set_col_widths(meta, [1.8, 4.7])

    doc.add_paragraph()

    # ── Dimensions table ─────────────────────────────────────────────────────
    h1 = doc.add_heading("Dimension Scores", level=2)
    h1.runs[0].font.color.rgb = rgb(BLUE)

    sorted_dims = sorted(
        report.dimensions.items(),
        key=lambda kv: -(report.weights_used.get(kv[0]) or 0),
    )
    dim_tbl = doc.add_table(rows=len(sorted_dims) + 1, cols=4)
    dim_tbl.style = "Table Grid"
    for cell, hdr in zip(dim_tbl.rows[0].cells, ["Dimension", "What it measures", "Score", "Weight"]):
        shade_cell(cell, NAVY)
        p = cell.paragraphs[0]
        r2 = p.add_run(hdr)
        r2.bold = True
        r2.font.color.rgb = rgb(WHITE)

    for (key, val), row in zip(sorted_dims, dim_tbl.rows[1:]):
        label, description = _DIM_LABELS.get(key, (key, ""))
        weight = report.weights_used.get(key)
        val_s = "—" if val is None else f"{val:.3f}"
        w_s = "excl." if weight is None else f"{weight:.2f}"
        dim_color = (GREEN if val is not None and val >= 0.90
                     else AMBER if val is not None and val >= 0.75
                     else RED if val is not None else None)
        for cell, text in zip(row.cells, [label, description, val_s, w_s]):
            p = cell.paragraphs[0]
            rn = p.add_run(text)
            if text == val_s and dim_color:
                rn.font.color.rgb = rgb(dim_color)
                rn.bold = True
    set_col_widths(dim_tbl, [1.8, 3.2, 0.8, 0.8])

    doc.add_paragraph()

    # ── Run summary table ─────────────────────────────────────────────────────
    h2 = doc.add_heading("Per-Run Summary", level=2)
    h2.runs[0].font.color.rgb = rgb(BLUE)

    run_tbl = doc.add_table(rows=len(results) + 1, cols=4)
    run_tbl.style = "Table Grid"
    for cell, hdr in zip(run_tbl.rows[0].cells, ["Run", "Status", "Overall Risk", "Time (s)"]):
        shade_cell(cell, NAVY)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r3 = p.add_run(hdr)
        r3.bold = True
        r3.font.color.rgb = rgb(WHITE)

    risk_values = (report.details or {}).get("overall_risk_values") or []
    for r, row in zip(results, run_tbl.rows[1:]):
        status = r.status or "?"
        s_color = GREEN if status in ("complete", "cache_hit") else RED
        risk = risk_values[r.index] if r.index < len(risk_values) else "?"
        r_color = (GREEN if risk == "Low" else AMBER if risk == "Medium" else RED if risk in ("High", "Critical") else None)
        for cell, (val, col, bold) in zip(row.cells, [
            (str(r.index + 1), None, False),
            (status, s_color, True),
            (str(risk), r_color, True),
            (str(r.elapsed_sec), None, False),
        ]):
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            rn = p.add_run(val)
            rn.bold = bold
            if col:
                rn.font.color.rgb = rgb(col)
    set_col_widths(run_tbl, [0.6, 1.5, 1.5, 1.0])

    doc.add_paragraph()

    # ── Verdict ───────────────────────────────────────────────────────────────
    h3 = doc.add_heading("Summary", level=2)
    h3.runs[0].font.color.rgb = rgb(BLUE)

    verdict = doc.add_paragraph()
    v1 = verdict.add_run(f"Consistency score: {score:.3f} [{band}]")
    v1.bold = True
    v1.font.color.rgb = rgb(score_color)

    if score >= 0.90:
        note = doc.add_paragraph()
        note.add_run("Output is highly consistent across runs. The analyser is reliably recommending the same tests and risk level for this CR.").font.color.rgb = rgb(GREEN)
    elif score >= 0.75:
        note = doc.add_paragraph()
        note.add_run("Output is acceptably consistent but some dimensions show variance. Review dimension scores above to identify areas for improvement.").font.color.rgb = rgb(AMBER)
    else:
        note = doc.add_paragraph()
        note.add_run("Output consistency is below the acceptable threshold. The analyser is producing materially different results across runs for this CR.").font.color.rgb = rgb(RED)

    if report.notes:
        doc.add_paragraph()
        hn = doc.add_heading("Notes", level=2)
        hn.runs[0].font.color.rgb = rgb(BLUE)
        for n in report.notes:
            p = doc.add_paragraph(style="List Bullet")
            p.add_run(n)

    doc.save(out_path)
    logger.info("Word report written to %s", out_path)
